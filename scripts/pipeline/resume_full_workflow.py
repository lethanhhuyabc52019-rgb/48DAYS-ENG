import os
import sys
import io
import time
import openpyxl
import subprocess

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

import gdown

ROOT_DIR = r"D:\2.English\Tai lieu_ENG\Drive_Download"
RESUME_LOG = os.path.join(ROOT_DIR, "download_resume.log")
SHEET_FILE = r"C:\Users\Admin\.gemini\antigravity-ide\brain\4d98faf6-7d4d-4d70-8579-afc4f2a66e57\scratch\sheet.xlsx"

DOC_EXTS = ('.pdf', '.docx', '.doc', '.xlsx', '.xls', '.pptx', '.ppt', '.txt')
VIDEO_EXTS = ('.mp4', '.mkv', '.webm', '.avi', '.mov')

def log(msg):
    ts = time.strftime("[%Y-%m-%d %H:%M:%S]")
    line = f"{ts} {msg}"
    print(line, flush=True)
    try:
        with open(RESUME_LOG, "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except Exception:
        pass

log("==================================================")
log("BẮT ĐẦU RESUME WORKFLOW DOWNLOAD (TASK 1 & TASK 2)")
log("==================================================")

# Load data from Google Sheet
wb = openpyxl.load_workbook(SHEET_FILE, data_only=False)
s = wb['48 ngày lấy gốc']

days_info = {}
for row in range(8, 110):
    val_a = s.cell(row=row, column=1).value
    val_b = s.cell(row=row, column=2).value
    val_c = s.cell(row=row, column=3).value
    hl_c = s.cell(row=row, column=3).hyperlink
    url_c = hl_c.target if hl_c else None
    
    val_e = s.cell(row=row, column=5).value
    hl_e = s.cell(row=row, column=5).hyperlink
    url_e = hl_e.target if hl_e else None
    
    if val_a and 'NGÀY' in str(val_a).upper():
        num_str = str(val_a).strip().replace("NGÀY", "").strip()
        try:
            day_num = int(num_str)
        except ValueError:
            continue
        days_info[day_num] = {
            "day_name": f"NGÀY {day_num}",
            "topic": str(val_b).strip() if val_b else "",
            "doc_url": url_e,
            "video_url": url_c
        }

log(f"Đã load {len(days_info)} ngày từ file dữ liệu Google Sheets.")

# ==================================================
# TASK 1: DOWNLOAD TÀI LIỆU CÒN THIẾU
# ==================================================
log("\n--- TASK 1: TIẾP TỤC CÁC FILE TÀI LIỆU CHƯA DOWNLOAD ---")

task1_skipped = []
task1_downloaded = []
task1_failed = []

for day_num in range(1, 49):
    info = days_info.get(day_num, {})
    day_name = f"NGÀY {day_num}"
    day_dir = os.path.join(ROOT_DIR, day_name)
    os.makedirs(day_dir, exist_ok=True)
    
    # Check current docs in folder
    existing_docs = [f for f in os.listdir(day_dir) if f.lower().endswith(DOC_EXTS) and os.path.getsize(os.path.join(day_dir, f)) > 0]
    
    # Each day typically has 3 doc files (Bài thi, Bài học, Đáp án).
    if len(existing_docs) >= 3:
        log(f"[{day_name}] Đã đủ tài liệu ({len(existing_docs)} file) -> Bỏ qua.")
        task1_skipped.append((day_name, len(existing_docs)))
        continue
        
    doc_url = info.get("doc_url")
    if not doc_url:
        log(f"[{day_name}] Không có link tài liệu trong sheet -> Bỏ qua.")
        continue
        
    log(f"[{day_name}] Đang tải tài liệu (Hiện có: {len(existing_docs)} file) -> {doc_url}")
    
    success = False
    for attempt in range(1, 3):
        try:
            before_files = set(os.listdir(day_dir))
            gdown.download_folder(doc_url, output=day_dir, quiet=True)
            after_files = set(os.listdir(day_dir))
            new_files = list(after_files - before_files)
            
            # Verify docs exist and > 0 bytes
            current_docs = [f for f in os.listdir(day_dir) if f.lower().endswith(DOC_EXTS) and os.path.getsize(os.path.join(day_dir, f)) > 0]
            
            if len(current_docs) > len(existing_docs):
                log(f" -> [{day_name}] THÀNH CÔNG (Tải thêm {len(new_files)} file, tổng hiện có: {len(current_docs)} docs):")
                for nf in new_files:
                    fpath = os.path.join(day_dir, nf)
                    size_kb = round(os.path.getsize(fpath) / 1024, 1)
                    log(f"     + {nf} ({size_kb} KB)")
                task1_downloaded.append((day_name, len(current_docs)))
                success = True
                break
            else:
                log(f" -> [{day_name}] Không có file mới nào được tải.")
                success = True # nothing new to download
                break
                
        except Exception as e:
            err_str = str(e)
            log(f" -> [{day_name}] Lần thử {attempt} thất bại: {err_str[:120]}")
            if "rate limit" in err_str.lower() or "too many" in err_str.lower():
                log("   [Rate Limit] Chờ 10 giây trước khi thử lại...")
                time.sleep(10)
            else:
                time.sleep(2)
                
    if not success:
        task1_failed.append((day_name, doc_url, "Google Drive Rate Limit / Download Error"))
        
    # Polite sleep to avoid rate limiting
    time.sleep(2)

# ==================================================
# TASK 2: KIỂM TRA 5 VIDEO TỪNG BỊ LỖI
# ==================================================
log("\n--- TASK 2: KIỂM TRA 5 VIDEO CÒN THIẾU ---")

failed_videos = [
    (8, "https://www.youtube.com/watch?v=9XZGT-UBbHY"),
    (10, "https://www.youtube.com/watch?v=GzX0BK03-ys"),
    (11, "https://www.youtube.com/watch?v=nRo4RQ23Gxk"),
    (12, "https://www.youtube.com/watch?v=aW55_W04ddc"),
    (13, "https://www.youtube.com/watch?v=8i_lkyZOh9g")
]

task2_downloaded = []
task2_still_failed = []

for day_num, vurl in failed_videos:
    day_name = f"NGÀY {day_num}"
    day_dir = os.path.join(ROOT_DIR, day_name)
    os.makedirs(day_dir, exist_ok=True)
    
    # Check if video already exists
    existing_vids = [f for f in os.listdir(day_dir) if any(f.lower().endswith(ext) for ext in VIDEO_EXTS) and os.path.getsize(os.path.join(day_dir, f)) > 0]
    if existing_vids:
        log(f"[{day_name}] Video đã tồn tại: {existing_vids[0]} -> Bỏ qua.")
        continue
        
    log(f"[{day_name}] Kiểm tra video: {vurl}")
    cmd = [
        "yt-dlp",
        "--no-playlist",
        "--retries", "2",
        "--no-part",
        "-P", day_dir,
        "-o", "%(title)s.%(ext)s",
        vurl
    ]
    p = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='replace')
    current_vids = [f for f in os.listdir(day_dir) if any(f.lower().endswith(ext) for ext in VIDEO_EXTS) and os.path.getsize(os.path.join(day_dir, f)) > 0]
    
    if current_vids:
        log(f" -> [{day_name}] TẢI VIDEO THÀNH CÔNG: {current_vids[0]}")
        task2_downloaded.append((day_name, current_vids[0]))
    else:
        err_msg = p.stderr.strip() or p.stdout.strip() or "Unavailable"
        last_line = err_msg.splitlines()[-1] if err_msg else "Unavailable"
        log(f" -> [{day_name}] Video vẫn không khả dụng: {last_line}")
        task2_still_failed.append((day_name, vurl, last_line[:120]))

# ==================================================
# FINAL RECONCILIATION & AUDIT
# ==================================================
log("\n==================================================")
log("TỔNG KẾT TOÀN BỘ WORKFLOW SAU KHI HOÀN TẤT")
log("==================================================")

final_docs_present = []
final_docs_missing = []
final_vids_present = []
final_vids_missing = []

for day_num in range(1, 49):
    day_name = f"NGÀY {day_num}"
    day_dir = os.path.join(ROOT_DIR, day_name)
    
    files = os.listdir(day_dir) if os.path.exists(day_dir) else []
    docs = [f for f in files if f.lower().endswith(DOC_EXTS) and os.path.getsize(os.path.join(day_dir, f)) > 0]
    vids = [f for f in files if any(f.lower().endswith(ext) for ext in VIDEO_EXTS) and os.path.getsize(os.path.join(day_dir, f)) > 0]
    
    if docs:
        final_docs_present.append((day_name, len(docs)))
    else:
        final_docs_missing.append(day_name)
        
    if vids:
        final_vids_present.append((day_name, len(vids)))
    else:
        final_vids_missing.append(day_name)

log(f"- Tổng ngày có tài liệu: {len(final_docs_present)}/48")
log(f"- Tổng ngày có video: {len(final_vids_present)}/48")
if final_docs_missing:
    log(f"- Tài liệu còn thiếu: {final_docs_missing}")
if final_vids_missing:
    log(f"- Video còn thiếu: {final_vids_missing}")
log("==================================================")
