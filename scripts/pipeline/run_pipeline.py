import os
import sys
import io
import re
import time
import shutil
import subprocess
import unicodedata
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8')

ROOT_DIR = r"D:\2.English\Tai lieu_ENG\Drive_Download"
PIPELINE_LOG = os.path.join(ROOT_DIR, "pipeline_execution.log")
SHEET_FILE = r"C:\Users\Admin\.gemini\antigravity-ide\brain\4d98faf6-7d4d-4d70-8579-afc4f2a66e57\scratch\sheet.xlsx"

def log(msg):
    ts = time.strftime("[%Y-%m-%d %H:%M:%S]")
    line = f"{ts} {msg}"
    print(line, flush=True)
    try:
        with open(PIPELINE_LOG, "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except Exception:
        pass

def strip_accents(text):
    text = unicodedata.normalize('NFD', text)
    text = ''.join(c for c in text if unicodedata.category(c) != 'Mn')
    text = text.replace('đ', 'd').replace('Đ', 'D')
    return re.sub(r'\s+', ' ', text).strip().lower()

# Day topics mapped to keywords
day_keywords = [
    (1, ["the khang dinh va phu dinh cua dong tu to be", "dong tu to be"]),
    (2, ["the nghi van cua dong tu to be"]),
    (3, ["cau hoi who va what", "cau hoi who va that"]),
    (4, ["cau hoi where va when"]),
    (5, ["dong tu thuong o hien tai", "the khang dinh cua dong tu thuong"]),
    (6, ["the phu dinh cua dong tu thuong"]),
    (7, ["the nghi van cua dong tu thuong"]),
    (8, ["thi hien tai don"]),
    (9, ["tu loai"]),
    (10, ["thi hien tai tiep dien"]),
    (11, ["phan biet thi hien tai don va hien tai tiep dien"]),
    (12, ["thi qua khu don the khang dinh", "dong tu khong dinh"]),
    (13, ["thi qua khu don the phu dinh va nghi van"]),
    (14, ["thi qua khu tiep dien"]),
    (15, ["thi hien tai hoan thanh"]),
    (16, ["thi tuong lai don"]),
    (17, ["thi tuong lai hoan thanh"]),
    (18, ["ngu am"]),
    (19, ["trong am"]),
    (20, ["tu de hoi"]),
    (21, ["so va ten"]),
    (22, ["dong tu khuyet thieu"]),
    (23, ["lien tu and so but because"]),
    (24, ["lien tu chi thoi gian"]),
    (25, ["lien tu chi su doi lap"]),
    (26, ["cau dieu kien loai 1"]),
    (27, ["cau dieu kien loai 2"]),
    (28, ["cau dieu kien loai 3"]),
    (29, ["luyen nghe dien tu"]),
    (30, ["luyen nghe chep chinh ta"]),
    (31, ["luyen nghe ve gio"]),
    (32, ["luyen nghe ngay thang"]),
    (33, ["luyen nghe dia diem"]),
    (34, ["luyen nghe ve tien bac"]),
    (35, ["dai tu phan than"]),
    (36, ["su hoa hop ve thi"]),
    (37, ["tieng anh giao tiep (1)", "tieng anh giao tiep 1"]),
    (38, ["tu tuong hinh"]),
    (39, ["cac quoc gia chau luc"]),
    (40, ["luyen nghe ve so thich"]),
    (41, ["phuong tien giao thong", "giao thong"]),
    (42, ["luyen nghe ve the thao", "the thao"]),
    (43, ["nghe nghiep"]),
    (44, ["cong nghe"]),
    (45, ["tieng anh giao tiep (2)", "tieng anh giao tiep 2"]),
    (46, ["note taking", "note-taking"]),
    (47, ["paraphrasing"]),
    (48, ["thuyet trinh", "gioi thieu ban than"])
]

def map_file_to_day(filename):
    if filename in ["download_log.txt", "pipeline_execution.log"]:
        return None
    s = strip_accents(filename)
    m = re.search(r'ngay\s*(\d+)', s)
    if m:
        return f"NGÀY {m.group(1)}"
        
    candidates = []
    for day_num, kws in day_keywords:
        for kw in kws:
            if kw in s:
                candidates.append((len(kw), day_num))
    if candidates:
        candidates.sort(key=lambda x: x[0], reverse=True)
        return f"NGÀY {candidates[0][1]}"
    return None

log("==================================================")
log("TASK 1: PHÂN LOẠI & MOVE FILE ĐÃ CÓ VÀO FOLDER NGÀY")
log("==================================================")

# 1. Scan files at root
root_entries = os.listdir(ROOT_DIR)
initial_files = [f for f in root_entries if os.path.isfile(os.path.join(ROOT_DIR, f))]
log(f"Tổng số file phát hiện tại root: {len(initial_files)}")

task1_moved = {}
task1_unidentified = []

for fn in initial_files:
    src_path = os.path.join(ROOT_DIR, fn)
    day_folder_name = map_file_to_day(fn)
    
    if day_folder_name:
        day_dir = os.path.join(ROOT_DIR, day_folder_name)
        os.makedirs(day_dir, exist_ok=True)
        dst_path = os.path.join(day_dir, fn)
        try:
            shutil.move(src_path, dst_path)
            task1_moved.setdefault(day_folder_name, []).append(fn)
            log(f" -> Moved: '{fn}' -> {day_folder_name}")
        except Exception as e:
            log(f" -> Error moving '{fn}': {e}")
            task1_unidentified.append(fn)
    else:
        task1_unidentified.append(fn)
        log(f" -> Giữ nguyên tại root (Chưa xác định / Log): '{fn}'")

# Validation
total_moved = sum(len(v) for v in task1_moved.values())
log(f"\n[Validation Task 1] Đã chuyển: {total_moved}, Chưa xác định: {len(task1_unidentified)}, Tổng ban đầu: {len(initial_files)}")
assert (total_moved + len(task1_unidentified)) == len(initial_files), "Mismatch in file counts!"
log("Validation Task 1 PASSED! Tự động chuyển sang Task 2 ngay lập tức.")

log("\n==================================================")
log("TASK 2: DOWNLOAD TUẦN TỰ TOÀN BỘ VIDEO")
log("==================================================")

# 1. Extract VIDEO links from sheet
wb = openpyxl.load_workbook(SHEET_FILE, data_only=False)
s = wb['48 ngày lấy gốc']

video_tasks = []
day_topics = {}
current_day = None

for row in range(8, 110):
    val_a = s.cell(row=row, column=1).value
    val_b = s.cell(row=row, column=2).value
    val_c = s.cell(row=row, column=3).value
    hl_c = s.cell(row=row, column=3).hyperlink
    target_c = hl_c.target if hl_c else None
    
    if val_a and 'NGÀY' in str(val_a).upper():
        current_day = str(val_a).strip()
        if val_b:
            day_topics[current_day] = str(val_b).strip()
            
    if val_c and 'VIDEO' in str(val_c).upper() and target_c:
        video_tasks.append({
            "day": current_day,
            "topic": day_topics.get(current_day, ""),
            "url": target_c,
            "row": row
        })

log(f"Tổng số link VIDEO tìm thấy trong Google Sheets: {len(video_tasks)}")

task2_downloaded = 0
task2_skipped = 0
task2_failed = []

VIDEO_EXTS = ['.mp4', '.mkv', '.webm', '.avi', '.mov']

for idx, v in enumerate(video_tasks, 1):
    day_name = v["day"]
    topic = v["topic"]
    url = v["url"]
    
    log(f"\n[{idx}/{len(video_tasks)}] Xử lý Video {day_name}: {topic}")
    log(f" -> URL: {url}")
    
    day_dir = os.path.join(ROOT_DIR, day_name)
    os.makedirs(day_dir, exist_ok=True)
    
    # Check if a completed video file already exists
    existing_videos = [
        f for f in os.listdir(day_dir)
        if any(f.lower().endswith(ext) for ext in VIDEO_EXTS) and os.path.getsize(os.path.join(day_dir, f)) > 0
    ]
    
    if existing_videos:
        log(f" -> [SKIP] Video đã tồn tại trong {day_name}: {existing_videos[0]} ({round(os.path.getsize(os.path.join(day_dir, existing_videos[0]))/1048576, 2)} MB)")
        task2_skipped += 1
        continue
        
    # Download video using yt-dlp
    cmd = [
        "yt-dlp",
        "--no-playlist",
        "--retries", "3",
        "--fragment-retries", "3",
        "--no-part",
        "-P", day_dir,
        "-o", "%(title)s.%(ext)s",
        url
    ]
    
    try:
        log(" -> Đang tải video...")
        p = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=300)
        
        # Verify video file exists and > 0 bytes
        current_videos = [
            f for f in os.listdir(day_dir)
            if any(f.lower().endswith(ext) for ext in VIDEO_EXTS) and os.path.getsize(os.path.join(day_dir, f)) > 0
        ]
        
        # Check no temporary files
        temp_files = [f for f in os.listdir(day_dir) if f.endswith(('.part', '.tmp', '.crdownload'))]
        
        if current_videos and not temp_files:
            vname = current_videos[0]
            vsize = round(os.path.getsize(os.path.join(day_dir, vname)) / 1048576, 2)
            log(f" -> [THÀNH CÔNG] Đã tải hoàn chỉnh: {vname} ({vsize} MB)")
            task2_downloaded += 1
        else:
            err_msg = p.stderr.strip() or p.stdout.strip() or "Unknown error"
            first_err = err_msg.splitlines()[-1] if err_msg else "No video downloaded"
            log(f" -> [LỖI] Tải thất bại: {first_err}")
            task2_failed.append({"day": day_name, "url": url, "error": first_err[:120]})
            
    except subprocess.TimeoutExpired:
        log(f" -> [LỖI] Quá thời gian tải (timeout 300s)")
        task2_failed.append({"day": day_name, "url": url, "error": "Timeout > 300s"})
    except Exception as e:
        log(f" -> [LỖI] Exception: {e}")
        task2_failed.append({"day": day_name, "url": url, "error": str(e)[:120]})

log("\n==================================================")
log("TỔNG KẾT TOÀN BỘ PIPELINE")
log("==================================================")
log("[BÁO CÁO TASK 1 - SẮP XẾP FILE]")
log(f"- Tổng số file đã quét: {len(initial_files)}")
log("- Đã chuyển vào folder:")
for d in sorted(task1_moved.keys(), key=lambda x: int(x.split()[1])):
    log(f"  + {d}: {len(task1_moved[d])} file")
log(f"- Chưa xác định (giữ ở root): {len(task1_unidentified)} file: {task1_unidentified}")

log("\n[BÁO CÁO TASK 2 - DOWNLOAD VIDEO]")
log(f"- Tổng link VIDEO tìm thấy: {len(video_tasks)}")
log(f"- Đã tải thành công: {task2_downloaded} file")
log(f"- Đã bỏ qua do tồn tại sẵn: {task2_skipped} file")
log(f"- Tải thất bại: {len(task2_failed)} file")
for fail in task2_failed:
    log(f"  + {fail['day']} ({fail['url']}) - Lỗi: {fail['error']}")
log("==================================================")
